import openpyxl
from datetime import datetime
from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.contrib.auth.models import User
from openpyxl.styles import Font, PatternFill, Alignment
from nucleo.models import ActividadLog, Inmueble, Factura, Pago
from nucleo.utils import sanitizar_para_excel

def es_administrador(user):
    return user.is_staff or user.is_superuser

@login_required
@user_passes_test(es_administrador)
def vista_auditoria(request):
    """Panel de auditoría global con filtros."""
    # Obtener todos los logs
    logs = ActividadLog.objects.select_related('usuario').all().order_by('-fecha_hora')
    
    # Filtros
    usuario_id = request.GET.get('usuario')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    modelo = request.GET.get('modelo')
    accion = request.GET.get('accion')
    
    # Aplicar filtros
    if usuario_id:
        logs = logs.filter(usuario_id=usuario_id)
    
    if fecha_desde:
        logs = logs.filter(fecha_hora__gte=fecha_desde)
    
    if fecha_hasta:
        logs = logs.filter(fecha_hora__lte=fecha_hasta)
    
    if modelo:
        logs = logs.filter(modelo=modelo)
    
    if accion:
        logs = logs.filter(accion=accion)
    
    # Paginación
    paginator = Paginator(logs, 50)
    page_number = request.GET.get('page')
    logs_page = paginator.get_page(page_number)
    
    # Opciones para filtros
    usuarios = User.objects.filter(is_active=True).order_by('username')
    modelos = ActividadLog.objects.values_list('modelo', flat=True).distinct().order_by('modelo')
    acciones = ActividadLog.ACCION_CHOICES
    
    return render(request, 'nucleo/auditoria/panel.html', {
        'logs': logs_page,
        'usuarios': usuarios,
        'modelos': modelos,
        'acciones': acciones,
        'filtros': {
            'usuario': usuario_id,
            'fecha_desde': fecha_desde,
            'fecha_hasta': fecha_hasta,
            'modelo': modelo,
            'accion': accion,
        }
    })


@login_required
@user_passes_test(es_administrador)
def exportar_auditoria_excel(request):
    """Exporta los logs de auditoría a Excel."""
    logs = ActividadLog.objects.select_related('usuario').all().order_by('-fecha_hora')
    
    usuario_id = request.GET.get('usuario')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')
    modelo = request.GET.get('modelo')
    accion = request.GET.get('accion')
    
    if usuario_id:
        logs = logs.filter(usuario_id=usuario_id)
    if fecha_desde:
        logs = logs.filter(fecha_hora__gte=fecha_desde)
    if fecha_hasta:
        logs = logs.filter(fecha_hora__lte=fecha_hasta)
    if modelo:
        logs = logs.filter(modelo=modelo)
    if accion:
        logs = logs.filter(accion=accion)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Fecha/Hora', 'Usuario', 'Acción', 'Modelo', 'Objeto ID', 'Descripción']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    for row, log in enumerate(logs, 2):
        ws.cell(row=row, column=1, value=log.fecha_hora.strftime('%d/%m/%Y %H:%M:%S'))
        ws.cell(row=row, column=2, value=sanitizar_para_excel(log.usuario.username))
        ws.cell(row=row, column=3, value=sanitizar_para_excel(log.get_accion_display()))
        ws.cell(row=row, column=4, value=sanitizar_para_excel(log.modelo))
        ws.cell(row=row, column=5, value=log.objeto_id)
        ws.cell(row=row, column=6, value=sanitizar_para_excel(log.descripcion))
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 50
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=auditoria_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_administrador)
def historial_inmueble(request, inmueble_id):
    """Muestra el historial de cambios de un inmueble."""
    inmueble = get_object_or_404(Inmueble, pk=inmueble_id)
    logs = ActividadLog.objects.filter(
        modelo='Inmueble',
        objeto_id=inmueble_id
    ).select_related('usuario').order_by('-fecha_hora')
    
    return render(request, 'nucleo/auditoria/historial.html', {
        'objeto': inmueble,
        'tipo': 'Inmueble',
        'logs': logs,
    })


@login_required
@user_passes_test(es_administrador)
def historial_factura(request, factura_id):
    """Muestra el historial de cambios de una factura."""
    factura = get_object_or_404(Factura, pk=factura_id)
    logs = ActividadLog.objects.filter(
        modelo='Factura',
        objeto_id=factura_id
    ).select_related('usuario').order_by('-fecha_hora')
    
    return render(request, 'nucleo/auditoria/historial.html', {
        'objeto': factura,
        'tipo': 'Factura',
        'logs': logs,
    })


@login_required
@user_passes_test(es_administrador)
def historial_pago(request, pago_id):
    """Muestra el historial de cambios de un pago."""
    pago = get_object_or_404(Pago, pk=pago_id)
    logs = ActividadLog.objects.filter(
        modelo='Pago',
        objeto_id=pago_id
    ).select_related('usuario').order_by('-fecha_hora')
    
    return render(request, 'nucleo/auditoria/historial.html', {
        'objeto': pago,
        'tipo': 'Pago',
        'logs': logs,
    })
