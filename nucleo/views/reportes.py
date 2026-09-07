import subprocess
import os
import tempfile
import uuid
import openpyxl
from datetime import datetime, date
from decimal import Decimal
from urllib.parse import urlparse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.db.models import Sum, Q
from django.contrib import messages
from weasyprint import HTML
from openpyxl.styles import Font, PatternFill, Alignment

from nucleo.models import Inmueble, Factura, Pago, Gasto, DatosCondominio, TokenAccesoTemporal
from nucleo.utils import (
    validar_alicuotas, 
    calcular_deuda_inmueble, 
    registrar_actividad,
    sanitizar_para_excel
)

def es_directiva_o_admin(user):
    """Verifica si el usuario es staff, superusuario o pertenece al grupo Directiva."""
    return user.is_staff or user.is_superuser or user.groups.filter(name='Directiva').exists()

# =============================================================================
# RECIBO POR INMUEBLE
# =============================================================================

def recibo_inmueble(request, factura_id, inmueble_id):
    """
    Muestra el recibo de un inmueble para una factura.
    
    Permite acceso de dos formas:
    1. Usuario autenticado (@login_required)
    2. Token temporal válido (para Puppeteer)
    """
    # Verificar autenticación
    autenticado = False
    
    # Opción 1: Usuario logeado
    if request.user.is_authenticated:
        autenticado = True
    
    # Opción 2: Token válido
    token_param = request.GET.get('token')
    if token_param and not autenticado:
        try:
            token = TokenAccesoTemporal.objects.get(token=token_param)
            if token.es_valido() and token.vista_destino == 'recibo_inmueble':
                # Verificar que los parámetros coincidan
                if (token.parametros.get('factura_id') == factura_id and 
                    token.parametros.get('inmueble_id') == inmueble_id):
                    # Token válido: marcar como usado
                    token.usado = True
                    token.save()
                    autenticado = True
        except TokenAccesoTemporal.DoesNotExist:
            pass
    
    # Si no está autenticado de ninguna forma, redirigir a login
    if not autenticado:
        from django.contrib.auth.views import redirect_to_login
        return redirect_to_login(request.get_full_path())
    
    # Usuario autenticado o token válido: mostrar recibo
    factura = get_object_or_404(Factura, id=factura_id)
    inmueble = get_object_or_404(Inmueble, id=inmueble_id)
    
    # Calcular deuda usando la función de utils
    resultado = calcular_deuda_inmueble(inmueble, factura)
    
    # Obtener datos del condominio
    datos_condominio = DatosCondominio.objects.first()
    tasa_cambio = datos_condominio.tasa_cambio_dolar if datos_condominio else None
    
    return render(request, 'nucleo/recibos/recibo.html', {
        'factura': factura,
        'inmueble': inmueble,
        'resultado': resultado,
        'datos_condominio': datos_condominio,
        'tasa_cambio': tasa_cambio,
    })


# =============================================================================
# GENERACIÓN DE PDFs SIMPLES (WEASYPRINT)
# =============================================================================

@login_required
def pdf_reporte_inmuebles(request):
    """Genera un PDF simple con el reporte de inmuebles."""
    inmuebles = Inmueble.objects.all().order_by('numero_apto')
    validacion = validar_alicuotas()
    
    context = {
        'inmuebles': inmuebles,
        'total_inmuebles': inmuebles.count(),
        'suma_alicuotas': validacion['suma_total'],
        'validacion_alicuotas': validacion['es_valida'],
        'inmuebles_con_alicuota': validacion['inmuebles_con_alicuota'],
        'inmuebles_sin_alicuota': len(validacion['inmuebles_sin_alicuota']),
        'fecha_generacion': datetime.now(),
    }
    
    html_string = render_to_string('nucleo/pdf_reporte_inmuebles.html', context, request=request)
    html = HTML(string=html_string)
    pdf_file = html.write_pdf()
    
    registrar_actividad(
        usuario=request.user,
        accion='GENERAR',
        modelo='PDF',
        objeto_id=0,
        descripcion='Generado PDF de reporte de inmuebles'
    )
    
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_inmuebles.pdf"'
    
    return response


# =============================================================================
# GENERACIÓN DE PDFs CON PUPPETEER (MOTOR PRINCIPAL)
# =============================================================================

@login_required
def pdf_recibo_puppeteer(request, factura_id, inmueble_id):
    """Genera un PDF del recibo usando Puppeteer."""
    factura = get_object_or_404(Factura, id=factura_id)
    inmueble = get_object_or_404(Inmueble, id=inmueble_id)
    
    token = TokenAccesoTemporal.objects.create(
        usuario=request.user,
        vista_destino='recibo_inmueble',
        parametros={
            'factura_id': factura_id,
            'inmueble_id': inmueble_id
        }
    )
    
    recibo_path = f'/recibos/factura/{factura_id}/inmueble/{inmueble_id}/?token={token.token}'
    recibo_url = request.build_absolute_uri(recibo_path)
    
    # Validar que el destino apunte a un host y protocolo seguro permitido
    parsed_url = urlparse(recibo_url)
    if parsed_url.scheme not in ('http', 'https') or (settings.ALLOWED_HOSTS and parsed_url.netloc not in settings.ALLOWED_HOSTS and parsed_url.netloc not in ('127.0.0.1:8000', 'localhost:8000', '127.0.0.1', 'localhost')):
        recibo_url = f"http://127.0.0.1:8000{recibo_path}"
    
    temp_dir = tempfile.gettempdir()
    pdf_filename = f'recibo_{inmueble.numero_apto}_{factura.periodo}_{uuid.uuid4().hex[:8]}.pdf'
    pdf_path = os.path.join(temp_dir, pdf_filename)
    
    script_path = os.path.join(settings.BASE_DIR, 'scripts', 'generate_pdf.js')
    
    try:
        result = subprocess.run(
            ['node', script_path, recibo_url, pdf_path],
            capture_output=True,
            text=True,
            timeout=30
        )
        
        if result.returncode != 0:
            messages.error(request, f'Error al generar PDF: {result.stderr}')
            return redirect('nucleo:recibo_inmueble', factura_id=factura_id, inmueble_id=inmueble_id)
        
        with open(pdf_path, 'rb') as pdf_file:
            pdf_content = pdf_file.read()
        
        os.remove(pdf_path)
        
        registrar_actividad(
            usuario=request.user,
            accion='GENERAR',
            modelo='Recibo',
            objeto_id=inmueble.id,
            descripcion=f'Generado PDF de recibo para {inmueble.numero_apto} - {factura.periodo}'
        )
        
        response = HttpResponse(pdf_content, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="recibo_{inmueble.numero_apto}_{factura.periodo}.pdf"'
        
        return response
        
    except subprocess.TimeoutExpired:
        messages.error(request, 'Timeout al generar PDF. Intente nuevamente.')
        return redirect('nucleo:recibo_inmueble', factura_id=factura_id, inmueble_id=inmueble_id)
    except Exception as e:
        messages.error(request, f'Error inesperado: {str(e)}')
        return redirect('nucleo:recibo_inmueble', factura_id=factura_id, inmueble_id=inmueble_id)


# =============================================================================
# EXPORTACIONES A EXCEL
# =============================================================================

@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_solvencia_excel(request, factura_id):
    """Exporta el reporte de solvencia de una factura a Excel."""
    factura = get_object_or_404(Factura, id=factura_id)
    datos_condominio = DatosCondominio.objects.first()
    tasa_actual = datos_condominio.tasa_cambio_dolar if datos_condominio else Decimal('1.0000')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Solvencia {factura.periodo}"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Apto', 'Propietario', 'Total a Pagar (Bs)', 'Pagado (Bs)', 'Deuda (Bs)', 'Ref Deuda ($)', 'Estado']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        
    inmuebles = Inmueble.objects.all().order_by('numero_apto')
    for row, inmueble in enumerate(inmuebles, 2):
        resultado = calcular_deuda_inmueble(inmueble, factura)
        deuda = resultado['saldo']
        deuda_usd = deuda / tasa_actual if tasa_actual else 0
        estado = 'SOLVENTE' if deuda <= 0 else 'DEUDOR'
        
        ws.cell(row=row, column=1, value=sanitizar_para_excel(inmueble.numero_apto))
        ws.cell(row=row, column=2, value=sanitizar_para_excel(inmueble.propietario))
        ws.cell(row=row, column=3, value=round(float(resultado['total_a_pagar']), 2))
        ws.cell(row=row, column=4, value=round(float(resultado['total_pagado']), 2))
        ws.cell(row=row, column=5, value=round(float(deuda), 2))
        ws.cell(row=row, column=6, value=round(float(deuda_usd), 2))
        
        cell_estado = ws.cell(row=row, column=7, value=sanitizar_para_excel(estado))
        cell_estado.font = Font(color="008000" if estado == 'SOLVENTE' else "FF0000", bold=True)
            
    for col, width in [('A', 10), ('B', 30), ('C', 15), ('D', 15), ('E', 15), ('F', 15), ('G', 15)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f'solvencia_{factura.periodo}_{datetime.now().strftime("%Y%m%d")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_recibos_factura(request, factura_id):
    """Exporta resumen de todos los recibos de una factura a Excel."""
    factura = get_object_or_404(Factura, id=factura_id)
    datos_condominio = DatosCondominio.objects.first()
    tasa_actual = datos_condominio.tasa_cambio_dolar if datos_condominio else Decimal('1.0000')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Recibos {factura.periodo}"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Apto', 'Propietario', 'Alícuota', 'G. Comunes (Bs)', 'G. No Comunes (Bs)', 
               'Total a Pagar (Bs)', 'Total Pagado (Bs)', 'Deuda (Bs)', 'Ref. Deuda ($)', 'Estado']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        
    try:
        año, mes = factura.periodo.split('-')
        if int(mes) == 12:
            fecha_limite = date(int(año) + 1, 1, 1)
        else:
            fecha_limite = date(int(año), int(mes) + 1, 1)
        inmuebles = Inmueble.objects.filter(
            Q(fecha_ingreso__isnull=True) | Q(fecha_ingreso__lt=fecha_limite)
        ).order_by('numero_apto')
    except (ValueError, AttributeError):
        inmuebles = Inmueble.objects.all().order_by('numero_apto')
    
    for row, inmueble in enumerate(inmuebles, 2):
        resultado = calcular_deuda_inmueble(inmueble, factura)
        deuda = resultado['saldo']
        deuda_usd = deuda / tasa_actual if tasa_actual else 0
        estado = 'SOLVENTE' if deuda <= 0 else 'DEUDOR'
        
        ws.cell(row=row, column=1, value=sanitizar_para_excel(inmueble.numero_apto))
        ws.cell(row=row, column=2, value=sanitizar_para_excel(inmueble.propietario or 'Sin asignar'))
        ws.cell(row=row, column=3, value=round(float(inmueble.alicuota or 0), 4))
        ws.cell(row=row, column=4, value=round(float(resultado.get('gasto_comun', 0)), 2))
        ws.cell(row=row, column=5, value=round(float(resultado.get('gasto_no_comun', 0)), 2))
        ws.cell(row=row, column=6, value=round(float(resultado['total_a_pagar']), 2))
        ws.cell(row=row, column=7, value=round(float(resultado['total_pagado']), 2))
        ws.cell(row=row, column=8, value=round(float(deuda), 2))
        ws.cell(row=row, column=9, value=round(float(deuda_usd), 2))
        
        cell_estado = ws.cell(row=row, column=10, value=sanitizar_para_excel(estado))
        cell_estado.font = Font(color="008000" if estado == 'SOLVENTE' else "FF0000", bold=True)
            
    for col, width in [('A', 12), ('B', 25), ('C', 12), ('D', 15), ('E', 18), ('F', 18), ('G', 18), ('H', 15), ('I', 15), ('J', 12)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f'recibos_{factura.periodo}_{datetime.now().strftime("%Y%m%d")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_recibos_todos(request):
    """Exporta resumen de todos los recibos de TODAS las facturas a Excel."""
    datos_condominio = DatosCondominio.objects.first()
    tasa_actual = datos_condominio.tasa_cambio_dolar if datos_condominio else Decimal('1.0000')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recibos Todas Facturas"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Factura', 'Apto', 'Propietario', 'Alícuota', 'Total a Pagar (Bs)', 
               'Total Pagado (Bs)', 'Deuda (Bs)', 'Ref. Deuda ($)', 'Estado']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    row_num = 2
    for factura in Factura.objects.all().order_by('-periodo'):
        try:
            año, mes = factura.periodo.split('-')
            if int(mes) == 12:
                fecha_limite = date(int(año) + 1, 1, 1)
            else:
                fecha_limite = date(int(año), int(mes) + 1, 1)
            inmuebles = Inmueble.objects.filter(
                Q(fecha_ingreso__isnull=True) | Q(fecha_ingreso__lt=fecha_limite)
            ).order_by('numero_apto')
        except (ValueError, AttributeError):
            inmuebles = Inmueble.objects.all().order_by('numero_apto')
        
        for inmueble in inmuebles:
            resultado = calcular_deuda_inmueble(inmueble, factura)
            deuda = resultado['saldo']
            deuda_usd = deuda / tasa_actual if tasa_actual else 0
            estado = 'SOLVENTE' if deuda <= 0 else 'DEUDOR'
            
            ws.cell(row=row_num, column=1, value=sanitizar_para_excel(factura.periodo))
            ws.cell(row=row_num, column=2, value=sanitizar_para_excel(inmueble.numero_apto))
            ws.cell(row=row_num, column=3, value=sanitizar_para_excel(inmueble.propietario or 'Sin asignar'))
            ws.cell(row=row_num, column=4, value=round(float(inmueble.alicuota or 0), 4))
            ws.cell(row=row_num, column=5, value=round(float(resultado['total_a_pagar']), 2))
            ws.cell(row=row_num, column=6, value=round(float(resultado['total_pagado']), 2))
            ws.cell(row=row_num, column=7, value=round(float(deuda), 2))
            ws.cell(row=row_num, column=8, value=round(float(deuda_usd), 2))
            cell_estado = ws.cell(row=row_num, column=9, value=sanitizar_para_excel(estado))
            cell_estado.font = Font(color="008000" if estado == 'SOLVENTE' else "FF0000", bold=True)
            row_num += 1
    
    for col, width in [('A', 12), ('B', 12), ('C', 25), ('D', 12), ('E', 18), ('F', 18), ('G', 15), ('H', 15), ('I', 12)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="recibos_todas_facturas_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_solvencia_todos(request):
    """Exporta la solvencia de TODAS las facturas a Excel."""
    datos_condominio = DatosCondominio.objects.first()
    tasa_actual = datos_condominio.tasa_cambio_dolar if datos_condominio else Decimal('1.0000')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solvencia Todas Facturas"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Factura', 'Apto', 'Propietario', 'A Pagar (Bs)', 'Pagado (Bs)', 'Deuda (Bs)', 'Ref. $ (USD)', 'Estado']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    row_num = 2
    for factura in Factura.objects.all().order_by('-periodo'):
        try:
            año, mes = factura.periodo.split('-')
            if int(mes) == 12:
                fecha_limite = date(int(año) + 1, 1, 1)
            else:
                fecha_limite = date(int(año), int(mes) + 1, 1)
            inmuebles = Inmueble.objects.filter(
                Q(fecha_ingreso__isnull=True) | Q(fecha_ingreso__lt=fecha_limite)
            ).order_by('numero_apto')
        except (ValueError, AttributeError):
            inmuebles = Inmueble.objects.all().order_by('numero_apto')
        
        for inmueble in inmuebles:
            resultado = calcular_deuda_inmueble(inmueble, factura)
            deuda = resultado['saldo']
            deuda_usd = deuda / tasa_actual if tasa_actual else 0
            estado = 'SOLVENTE' if deuda <= 0 else 'DEUDOR'
            
            ws.cell(row=row_num, column=1, value=sanitizar_para_excel(factura.periodo))
            ws.cell(row=row_num, column=2, value=sanitizar_para_excel(inmueble.numero_apto))
            ws.cell(row=row_num, column=3, value=sanitizar_para_excel(inmueble.propietario or 'Sin asignar'))
            ws.cell(row=row_num, column=4, value=round(float(resultado['total_a_pagar']), 2))
            ws.cell(row=row_num, column=5, value=round(float(resultado['total_pagado']), 2))
            ws.cell(row=row_num, column=6, value=round(float(deuda), 2))
            ws.cell(row=row_num, column=7, value=round(float(deuda_usd), 2))
            cell_estado = ws.cell(row=row_num, column=8, value=sanitizar_para_excel(estado))
            cell_estado.font = Font(color="008000" if estado == 'SOLVENTE' else "FF0000", bold=True)
            row_num += 1
    
    for col, width in [('A', 12), ('B', 12), ('C', 25), ('D', 15), ('E', 15), ('F', 15), ('G', 15), ('H', 12)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="solvencia_todas_facturas_{datetime.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_inmuebles_excel(request):
    """Exporta la lista de inmuebles a Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inmuebles"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal='center', vertical='center')
    
    headers = ['Número', 'Propietario', 'Documento', 'Email', 'Teléfono', 'Alícuota (%)', 'Notas']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    
    inmuebles = Inmueble.objects.all().order_by('numero_apto')
    for row, inmueble in enumerate(inmuebles, 2):
        ws.cell(row=row, column=1, value=sanitizar_para_excel(inmueble.numero_apto))
        ws.cell(row=row, column=2, value=sanitizar_para_excel(inmueble.propietario or ''))
        ws.cell(row=row, column=3, value=sanitizar_para_excel(inmueble.documento_identidad or ''))
        ws.cell(row=row, column=4, value=sanitizar_para_excel(inmueble.email or ''))
        ws.cell(row=row, column=5, value=sanitizar_para_excel(inmueble.telefono or ''))
        ws.cell(row=row, column=6, value=round(float(inmueble.alicuota), 4) * 100)
        ws.cell(row=row, column=7, value=sanitizar_para_excel(inmueble.notas or ''))
    
    for col, width in [('A', 12), ('B', 25), ('C', 15), ('D', 30), ('E', 15), ('F', 12), ('G', 40)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=inmuebles_{datetime.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_facturas_excel(request):
    """Exporta la lista de facturas a Excel."""
    periodo_filtro = request.GET.get('periodo', None)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Período', 'Monto Total (Bs)', '# Gastos', 'Fecha Creación']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    if periodo_filtro:
        facturas = Factura.objects.filter(periodo=periodo_filtro).order_by('-periodo')
        filename_suffix = periodo_filtro
    else:
        facturas = Factura.objects.all().order_by('-periodo')
        filename_suffix = 'todos'
    
    for row, factura in enumerate(facturas, 2):
        monto_total = factura.gastos.aggregate(total=Sum('monto'))['total'] or Decimal('0.00')
        ws.cell(row=row, column=1, value=sanitizar_para_excel(factura.periodo))
        ws.cell(row=row, column=2, value=float(monto_total))
        ws.cell(row=row, column=3, value=factura.gastos.count())
        ws.cell(row=row, column=4, value=factura.creada_en.strftime('%d/%m/%Y'))
    
    for col, width in [('A', 15), ('B', 18), ('C', 12), ('D', 18)]:
        ws.column_dimensions[col].width = width
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=facturas_{filename_suffix}_{datetime.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_pagos_excel(request):
    """Exporta la lista de pagos a Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pagos"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    headers = ['Fecha', 'Inmueble', 'Factura', 'Monto', 'Moneda', 'Tasa', 'Total Bs', 'Método']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    
    pagos = Pago.objects.select_related('inmueble', 'factura').all().order_by('-fecha_pago')
    total_recaudado = Decimal('0.00')
    
    for row, pago in enumerate(pagos, 2):
        if pago.moneda == 'USD' and pago.tasa_cambio:
            total_bs = pago.monto_pagado * pago.tasa_cambio
        else:
            total_bs = pago.monto_pagado
        total_recaudado += total_bs
        
        ws.cell(row=row, column=1, value=pago.fecha_pago.strftime('%d/%m/%Y'))
        ws.cell(row=row, column=2, value=sanitizar_para_excel(str(pago.inmueble)))
        ws.cell(row=row, column=3, value=sanitizar_para_excel(str(pago.factura.periodo)))
        ws.cell(row=row, column=4, value=float(pago.monto_pagado))
        ws.cell(row=row, column=5, value=sanitizar_para_excel(pago.moneda))
        ws.cell(row=row, column=6, value=float(pago.tasa_cambio) if pago.tasa_cambio else '')
        ws.cell(row=row, column=7, value=float(total_bs))
        ws.cell(row=row, column=8, value=sanitizar_para_excel(pago.get_metodo_pago_display()))
    
    total_row = len(pagos) + 2
    ws.cell(row=total_row, column=6, value='TOTAL RECAUDADO:').font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=float(total_recaudado)).font = Font(bold=True)
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=pagos_{datetime.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response


@login_required
@user_passes_test(es_directiva_o_admin)
def exportar_resumen_financiero_excel(request):
    """Exporta un resumen financiero completo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    
    header_fill = PatternFill(start_color="2C5F8D", end_color="2C5F8D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="2C5F8D")
    
    ws['A1'] = 'RESUMEN FINANCIERO - SAC WEB'
    ws['A1'].font = title_font
    ws.merge_cells('A1:B1')
    
    ws['A3'] = 'Concepto'
    ws['B3'] = 'Monto (Bs)'
    ws['A3'].fill = header_fill
    ws['B3'].fill = header_fill
    ws['A3'].font = header_font
    ws['B3'].font = header_font
    
    total_inmuebles = Inmueble.objects.count()
    total_facturas = Factura.objects.count()
    
    pagos = Pago.objects.all()
    total_recaudado = Decimal('0.00')
    for pago in pagos:
        if pago.moneda == 'USD' and pago.tasa_cambio:
            total_recaudado += pago.monto_pagado * pago.tasa_cambio
        else:
            total_recaudado += pago.monto_pagado
    
    total_gastos = Gasto.objects.aggregate(total=Sum('monto'))['total'] or Decimal('0.00')
    deuda_pendiente = total_gastos - total_recaudado
    
    ws['A4'] = 'Total Inmuebles'
    ws['B4'] = total_inmuebles
    ws['A5'] = 'Total Facturas'
    ws['B5'] = total_facturas
    ws['A6'] = 'Total Recaudado'
    ws['B6'] = float(total_recaudado)
    ws['A7'] = 'Total Gastos'
    ws['B7'] = float(total_gastos)
    ws['A8'] = 'Deuda Pendiente'
    ws['B8'] = float(deuda_pendiente)
    
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=resumen_{datetime.now().strftime("%Y%m%d")}.xlsx'
    wb.save(response)
    return response
